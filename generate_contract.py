#!/usr/bin/env python3
"""
LandCulture Investment Ltd — Contract of Sale Generator
========================================================
Usage:
  Chat mode (default):
      python3 generate_contract.py

  Batch mode:
      python3 generate_contract.py --batch customers.xlsx

  Optional output directory:
      python3 generate_contract.py --output-dir ./contracts
      python3 generate_contract.py --batch customers.xlsx --output-dir ./contracts
"""

import argparse
import os
import sys
from datetime import date

try:
    from docx import Document
except ImportError:
    print("ERROR: python-docx not installed. Run: pip install python-docx")
    sys.exit(1)

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

# ---------------------------------------------------------------------------
# Template path (always relative to this script's directory)
# ---------------------------------------------------------------------------
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(SCRIPT_DIR, "contract_template.docx")

# ---------------------------------------------------------------------------
# Field definitions — (key, description, example_value)
# These map directly to {{KEY}} placeholders in the template.
# ---------------------------------------------------------------------------
FIELDS = [
    ("CONTRACT_DAY",            "Contract day (ordinal)",              "5th"),
    ("CONTRACT_MONTH",          "Contract month",                      "June"),
    ("CONTRACT_YEAR",           "Contract year",                       "2026"),
    ("CUSTOMER_NAME",           "Customer full name in CAPS",          "MR. OJOLOWO BOLUWATIFE"),
    ("CUSTOMER_ADDRESS",        "Customer address",                    "9, OYEBANKE OSUNSANYA CRESCENT, LAGOS"),
    ("NUM_PLOTS_WORDS",         "Number of plots — words+digits combo","Two (2)"),
    ("NUM_PLOTS_DIGITS",        "Number of plots — digit only",        "2"),
    ("PLOT_SIZE_SQM",           "Plot size in sqm",                    "900"),
    ("TOTAL_PRICE_DIGITS",      "Total price in digits",               "3,000,000"),
    ("TOTAL_PRICE_WORDS",       "Total price in CAPS words",           "THREE MILLION NAIRA ONLY"),
    ("DEPOSIT_DIGITS",          "Deposit amount in digits",            "800,000.00"),
    ("DEPOSIT_WORDS",           "Deposit amount in CAPS words",        "EIGHT HUNDRED THOUSAND NAIRA ONLY"),
    ("BALANCE_DIGITS",          "Balance amount in digits",            "2,200,000.00"),
    ("BALANCE_WORDS",           "Balance amount in CAPS words",        "TWO MILLION, TWO HUNDRED THOUSAND NAIRA ONLY"),
    ("PAYMENT_START_DATE",      "Payment start date",                  "26th March, 2026"),
    ("PAYMENT_DURATION_MONTHS", "Payment duration in months",          "12"),
    ("PAYMENT_END_DATE",        "Payment end date",                    "26th Day of March, 2027"),
    ("PAYMENT_START_DAY_FULL",  "Payment deadline in full CAPS",       "26TH MARCH, 2027"),
]

# NUM_PLOTS_DIGITS_UPPER is auto-derived from NUM_PLOTS_WORDS (uppercased) — not asked separately
FIELD_KEYS = [f[0] for f in FIELDS] + ["NUM_PLOTS_DIGITS_UPPER"]


# ---------------------------------------------------------------------------
# Core replacement helpers
# ---------------------------------------------------------------------------

def _replace_in_paragraph(paragraph, old, new):
    """Replace *old* with *new* in a paragraph, handling text split across runs."""
    if old not in paragraph.text:
        return False

    # Pass 1: simple per-run replacement
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)

    # Pass 2: if still present the text is split across runs — rebuild
    if old in paragraph.text:
        full_text = "".join(run.text for run in paragraph.runs)
        if old in full_text:
            new_text = full_text.replace(old, new)
            if paragraph.runs:
                paragraph.runs[0].text = new_text
                for run in paragraph.runs[1:]:
                    run.text = ""

    return True


def _replace_in_doc(doc, old, new):
    """Replace *old* with *new* everywhere in *doc* (paragraphs + table cells)."""
    for para in doc.paragraphs:
        _replace_in_paragraph(para, old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_in_paragraph(para, old, new)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def generate_contract(data: dict, output_dir: str = ".") -> str:
    """
    Fill the contract template with *data* and write a .docx file.

    Parameters
    ----------
    data : dict
        Keys must include all entries in FIELD_KEYS.
    output_dir : str
        Directory where the output file will be saved.

    Returns
    -------
    str
        Absolute path of the generated file.
    """
    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(
            f"Template not found: {TEMPLATE_PATH}\n"
            "Run create_template.py first to generate the template."
        )

    doc = Document(TEMPLATE_PATH)

    # Auto-derive NUM_PLOTS_DIGITS_UPPER if not explicitly provided
    if "NUM_PLOTS_DIGITS_UPPER" not in data or not data["NUM_PLOTS_DIGITS_UPPER"]:
        data = dict(data)
        data["NUM_PLOTS_DIGITS_UPPER"] = data.get("NUM_PLOTS_WORDS", "").upper()

    for key in FIELD_KEYS:
        value = data.get(key, "")
        placeholder = "{{" + key + "}}"
        _replace_in_doc(doc, placeholder, str(value))

    # Build safe filename
    customer_safe = (
        data.get("CUSTOMER_NAME", "CUSTOMER")
        .replace(" ", "_")
        .replace(".", "")
        .replace("/", "_")
        .replace("\\", "_")
    )
    today = date.today().strftime("%Y-%m-%d")
    filename = f"CONTRACT_{customer_safe}_{today}.docx"

    os.makedirs(output_dir, exist_ok=True)
    out_path = os.path.join(os.path.abspath(output_dir), filename)
    doc.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# Chat (interactive) mode
# ---------------------------------------------------------------------------

def _collect_fields_chat() -> dict:
    """Prompt the user one question at a time and return the collected data dict."""
    print("\n=== LandCulture Investment Ltd — Contract of Sale Generator ===")
    print("Answer each question. Press Enter to use the example value shown in brackets.\n")

    data = {}
    for key, description, example in FIELDS:
        prompt = f"{description} [{example}]: "
        while True:
            value = input(prompt).strip()
            if not value:
                value = example
                print(f"  Using default: {value}")
            if value:
                data[key] = value
                break

    return data


def _print_summary(data: dict):
    print("\n--- Summary of entered values ---")
    for key, description, _ in FIELDS:
        print(f"  {key}: {data.get(key, '')}")
    print("---------------------------------\n")


def chat_mode(output_dir: str = "."):
    data = _collect_fields_chat()
    _print_summary(data)

    while True:
        confirm = input("Generate contract with these values? [y/n]: ").strip().lower()
        if confirm == "y":
            break
        elif confirm == "n":
            print("Aborted.")
            return
        else:
            print("Please enter 'y' or 'n'.")

    out_path = generate_contract(data, output_dir=output_dir)
    print(f"\nContract generated successfully: {out_path}")


# ---------------------------------------------------------------------------
# Batch mode
# ---------------------------------------------------------------------------

def batch_mode(xlsx_path: str, output_dir: str = "."):
    """Read customers from an Excel file and generate one contract per row."""
    if not os.path.exists(xlsx_path):
        print(f"ERROR: Excel file not found: {xlsx_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    print(f"Columns in Excel: {[h for h in headers if h]}")

    col_map = {h: i for i, h in enumerate(headers) if h}

    missing = [k for k in FIELD_KEYS if k not in col_map]
    if missing:
        print("WARNING: The following expected columns are missing from the Excel file:")
        for f in missing:
            print(f"  - {f}")

    generated = 0
    errors = 0
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None for v in row):
            continue  # skip empty rows

        data = {}
        for key in FIELD_KEYS:
            if key in col_map:
                val = row[col_map[key]]
                data[key] = str(val) if val is not None else ""
            else:
                data[key] = ""

        try:
            out_path = generate_contract(data, output_dir=output_dir)
            print(f"  Row {row_idx}: Generated {os.path.basename(out_path)}")
            generated += 1
        except Exception as e:
            print(f"  Row {row_idx}: ERROR — {e}")
            errors += 1

    print(f"\nDone. {generated} contract(s) generated, {errors} error(s).")


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="LandCulture Investment Ltd — Contract of Sale Generator"
    )
    parser.add_argument(
        "--batch",
        metavar="EXCEL_FILE",
        help="Path to customers Excel (.xlsx) file for batch generation",
    )
    parser.add_argument(
        "--output-dir",
        default=".",
        metavar="DIR",
        help="Directory for generated contracts (default: current directory)",
    )
    args = parser.parse_args()

    if args.batch:
        batch_mode(args.batch, output_dir=args.output_dir)
    else:
        chat_mode(output_dir=args.output_dir)


if __name__ == "__main__":
    main()
